"""Build synchronized bilingual Word reading manuscripts from current paired sources.

No model, experiment or evidence-generation calls. Native equations and tables are
produced by Pandoc, then styled with python-docx. Export PDFs with native Word and
inspect page renders before publishing. The reference DOCX supplies styles only.
"""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[2]
VENUE = ROOT / "paper/venues/ncs"
SCALE = 0.4564026975


def run(command: list[str]) -> None:
    result = subprocess.run(command, capture_output=True, text=True, encoding="utf-8")
    if result.returncode:
        raise RuntimeError(result.stdout + result.stderr)
    if result.stderr.strip():
        print(result.stderr, flush=True)


def author_block(meta: dict, language: str) -> str:
    """Render the identified manuscript's author metadata without changing its order."""
    authors = meta.get("author", [])
    if not authors:
        return ""
    names = []
    for author in authors:
        marks = [str(author["affiliation_markers"])]
        if author.get("equal_contribution"):
            marks.append("†")
        if author.get("corresponding"):
            marks.append(r"\*")
        names.append(f"{author['name']}^{','.join(marks)}^")
    text = '::: {custom-style="Manuscript Authors"}\n' + ", ".join(names) + "\n:::\n\n"
    for affiliation in meta.get("affiliation", []):
        text += (
            '::: {custom-style="Manuscript Affiliation"}\n'
            f"^{affiliation['id']}^ {affiliation['name']}\n:::\n\n"
        )
    if meta.get("equal_contribution_note"):
        note = meta.get("equal_contribution_note_zh", meta["equal_contribution_note"])
        if language == "en":
            note = meta["equal_contribution_note"]
        text += '::: {custom-style="Manuscript Author Note"}\n† ' + note + "\n:::\n\n"
    if meta.get("correspondence"):
        label = "Correspondence" if language == "en" else "通讯作者"
        text += (
            '::: {custom-style="Manuscript Author Note"}\n'
            f"\\* {label}: {meta['correspondence']}\n:::\n\n"
        )
    return text


def prepare_markdown(language: str, build: Path) -> tuple[Path, dict]:
    from pypdf import PdfReader

    main = VENUE / ("article.md" if language == "en" else "article_zh.md")
    supplementary = VENUE / f"supplementary_{language}.md"
    _, raw, body = main.read_text("utf-8").split("---", 2)
    meta = yaml.safe_load(raw)
    labels = (
        ("Abstract", "References", "Supplementary information")
        if language == "en"
        else ("摘要", "参考文献", "补充材料")
    )
    text = (
        f'::: {{custom-style="Title"}}\n{meta["title"]}\n:::\n\n'
        + author_block(meta, language)
        + (
            f"# {labels[0]}\n\n{meta['abstract'].strip()}\n\n{body.strip()}\n\n"
            f"# {labels[1]}\n\n::: {{#refs}}\n:::\n\n"
            f"# {labels[2]}\n\n{supplementary.read_text('utf-8')}"
        )
    )
    assets = build / "figure-images"
    assets.mkdir(exist_ok=True)

    def image(match: re.Match) -> str:
        caption, relative = match.groups()
        path = (VENUE / relative).resolve()
        page = PdfReader(path).pages[0]
        width, height = float(page.mediabox.width), float(page.mediabox.height)
        png = path.with_suffix(".png")
        if not png.exists():
            png = assets / (path.parent.name + "-" + path.stem + ".png")
            if not png.exists():
                # Rasterization preserves native PDF geometry; source artwork is unchanged.
                import pypdfium2

                with pypdfium2.PdfDocument(path) as pdf:
                    pdf[0].render(scale=3).to_pil().save(png)
        w, h = width * SCALE / 72, height * SCALE / 72
        assert w <= 6.67 and h <= 8, (path, w, h)
        return (
            f"![]({png.as_posix()}){{width={w:.7f}in height={h:.7f}in}}\n\n"
            f'::: {{custom-style="Caption"}}\n{caption}\n:::'
        )

    text = re.sub(r"!\[([^\n]+)\]\(([^)]+\.pdf)\)", image, text)
    # Table captions immediately after a table should stay with it.
    text = re.sub(
        r"(?m)^(Table [A-H]\d+\.[^\n]*(?:\n(?!\n|#|\|)[^\n]+)*)",
        r'::: {custom-style="Caption"}\n\1\n:::',
        text,
    )
    text = re.sub(r"(?m)^(表 [A-H]\d+\.[^\n]+)", r'::: {custom-style="Caption"}\n\1\n:::', text)
    out = build / f"manuscript-{language}.md"
    out.write_text(text, "utf-8")
    return out, meta


def style_docx(path: Path, language: str, meta: dict) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document(path)
    # Pandoc can carry unused images from the reference DOCX into each rebuild.
    # Keep every image relationship used by any element in the document body.
    relationship_namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    used_relationships = {
        value
        for element in doc.element.iter()
        for key, value in element.attrib.items()
        if key.startswith(relationship_namespace)
    }
    for relationship_id, relationship in list(doc.part.rels.items()):
        if relationship.reltype == RT.IMAGE and relationship_id not in used_relationships:
            doc.part.drop_rel(relationship_id)
    doc.core_properties.author = meta.get("pdf_author", "")
    doc.core_properties.title = meta["title"]
    for sec in doc.sections:
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        sec.left_margin = sec.right_margin = Cm(2.2)
        sec.footer_distance = Cm(1.1)
    for name, size, bold in [
        ("Normal", 11 if language == "en" else 12, False),
        ("Body Text", 11 if language == "en" else 12, False),
        ("Title", 18, True),
        ("Heading 1", 16, True),
        ("Heading 2", 13, True),
        ("Caption", 10 if language == "en" else 10.5, False),
        ("Bibliography", 10, False),
        ("Manuscript Authors", 11, False),
        ("Manuscript Affiliation", 9.5, False),
        ("Manuscript Author Note", 9.5, False),
    ]:
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ("ascii", "hAnsi", "cs"):
            fonts.set(qn("w:" + key), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "SimHei" if bold else "SimSun")
        for attr in list(fonts.attrib):
            if attr.endswith("Theme"):
                del fonts.attrib[attr]
        fmt = style.paragraph_format
        fmt.space_after = Pt(7)
        fmt.line_spacing = 1.25 if language == "en" else 1.35
        fmt.widow_control = True
        if name in ("Title", "Heading 1", "Heading 2"):
            fmt.keep_with_next = True
        if name == "Caption":
            fmt.line_spacing = 1.05
            fmt.space_after = Pt(9)
            fmt.keep_together = True
        if name == "Bibliography":
            fmt.line_spacing = 1.05
            fmt.space_after = Pt(6)
        if name.startswith("Manuscript "):
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.line_spacing = 1.1
            fmt.space_before = Pt(0)
            fmt.space_after = Pt(4)
            fmt.keep_with_next = True
            fmt.keep_together = True
    in_supp = False
    in_references = False
    appendix = ""
    for p in doc.paragraphs:
        t = p.text.strip()
        appendix_match = re.match(r"^(?:Appendix|附录) ([A-H])\.", t)
        if appendix_match:
            appendix = appendix_match[1]
        p.paragraph_format.widow_control = True
        if t in ("Supplementary information", "补充材料"):
            in_supp = True
            in_references = False
        if t in ("References", "参考文献"):
            in_references = True
        if p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if t in (
            "Discussion",
            "讨论",
            "Methods",
            "方法",
            "References",
            "参考文献",
            "Supplementary information",
            "补充材料",
        ) or t.startswith(("Appendix ", "附录 ")):
            p.paragraph_format.page_break_before = True
        if t.startswith(("Appendix A.", "附录 A.")):
            p.paragraph_format.page_break_before = False
        if in_references and p.style.name not in ("Heading 1", "Heading 2"):
            p.paragraph_format.line_spacing = 1.05
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_together = True
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            for r in p.runs:
                r.font.size = Pt(10)
        if p._p.xpath(".//w:drawing"):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(0)
        if p.style.name == "Caption":
            p.paragraph_format.keep_together = True
            # Bold the label only, retaining editable caption text.
            m = re.match(
                r"^(?:Supplementary Figure|Figure|补充图|图)\s+[S\d]+(?:[a-z\u2013,、]+)?[. ]", t
            )
            if m:
                p.clear()
                p.add_run(t[: m.end()]).bold = True
                p.add_run(t[m.end() :])
        if in_supp and p.style.name in ("Normal", "Body Text"):
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                if not r.font.superscript and not r.font.subscript:
                    r.font.size = Pt(10 if language == "en" else 11)
        # Compact repeated table headings and the final equation appendix,
        # preserving font sizes while preventing near-empty trailing pages.
        if appendix == "B" and p.style.name == "Heading 2":
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
        if appendix == "H":
            p.paragraph_format.space_after = Pt(4)
            if p.style.name == "Heading 2":
                p.paragraph_format.space_before = Pt(10)
            if p.style.name in ("Normal", "Body Text"):
                p.paragraph_format.line_spacing = 1.1
        if (
            t.startswith(("Table ", "Supplementary Table ", "表 ", "补充表 "))
            and p._p.getnext() is not None
            and p._p.getnext().tag == qn("w:tbl")
        ):
            p.paragraph_format.keep_with_next = True
    for table in doc.tables:
        table.autofit = False
        cols = len(table.columns)
        headers = [c.text for c in table.rows[0].cells]
        metric_table = headers[0] in ("Readout", "读出")
        widths = [1.0] * cols
        if cols == 8:
            widths = [0.55, 0.95, 0.35, 0.75, 0.75, 0.75, 0.75, 0.8]
            if headers[0] == "Q":
                widths = [0.4, 0.72, 0.5, 0.55, 0.7, 0.72, 1.15, 0.95]
        elif cols == 6 and headers[0] in ("Response", "响应"):
            widths = [1.1, 1.1, 0.85, 0.85, 0.85, 1.15]
        elif cols == 5 and headers[0] in ("Model", "模型"):
            widths = [1.3, 1.0, 1.1, 1.1, 1.1]
        widths = [Cm(16.6 * v / sum(widths)) for v in widths]
        for col, width in zip(table.columns, widths, strict=True):
            col.width = width
        after = table._tbl.getnext()
        after_is_caption = (
            after is not None
            and after.tag == qn("w:p")
            and bool(after.xpath('.//w:pStyle[@w:val="Caption"]'))
        )
        for ri, row in enumerate(table.rows):
            pr = row._tr.get_or_add_trPr()
            no_split = OxmlElement("w:cantSplit")
            pr.append(no_split)
            if ri == 0:
                pr.append(OxmlElement("w:tblHeader"))
            for ci, cell in enumerate(row.cells):
                cell.width = widths[ci]
                for p in cell.paragraphs:
                    p.paragraph_format.space_after = Pt(2 if metric_table else 3)
                    p.paragraph_format.space_before = Pt(2 if metric_table else 3)
                    p.paragraph_format.line_spacing = 1.05
                    p.paragraph_format.keep_with_next = ri < len(table.rows) - 1 or after_is_caption
                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(9 if language == "en" else 9.5)
                        if ri == 0:
                            r.bold = True
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for sec in doc.sections:
        footer = sec.footer.paragraphs[0]
        if not footer._p.xpath(".//w:fldSimple") and not footer._p.xpath(".//w:instrText"):
            footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
            field = OxmlElement("w:fldSimple")
            field.set(qn("w:instr"), "PAGE")
            footer._p.append(field)
    doc.save(path)
    print(
        f"{language}: {len(doc.inline_shapes)} figures; {len(doc.tables)} native tables; "
        f"{len(doc.element.xpath('.//m:oMath'))} native math objects",
        flush=True,
    )


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--build", type=Path, required=True)
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--dependency-path", type=Path)
    parser.add_argument("--language", choices=["en", "zh", "both"], default="both")
    args = parser.parse_args()
    if args.dependency_path:
        sys.path.insert(0, str(args.dependency_path))
    args.build.mkdir(parents=True, exist_ok=True)
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise RuntimeError("Pandoc is required")
    for language in ["en", "zh"] if args.language == "both" else [args.language]:
        source, meta = prepare_markdown(language, args.build)
        output = args.build / f"chemworld-ncs-{language}-final.docx"
        # All headings/captions are explicit translations; English CSL preserves
        # original-language bibliographic titles in both manuscript versions.
        run(
            [
                pandoc,
                str(source),
                "--from=markdown+tex_math_dollars+superscript",
                "--to=docx",
                "--reference-doc=" + str(args.reference),
                "--citeproc",
                "--bibliography=" + str(ROOT / "paper/chemworld_integrated_references.bib"),
                "--csl=" + str(VENUE / "numeric-references.csl"),
                "--metadata=lang:en-US",
                "--output=" + str(output),
            ]
        )
        style_docx(output, language, meta)
        print(f"Built {output}", flush=True)


if __name__ == "__main__":
    main()
