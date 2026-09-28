"""Export the Chinese manuscript, or append Methods without rebuilding author edits."""

from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "paper/venues/ncs/article_zh_main.md"
ENGLISH = SOURCE.with_name("article.md")
METHODS = SOURCE.with_name("article_zh_methods.md")


def figure_sizes_pt() -> dict[Path, tuple[float, float]]:
    """Native PDF dimensions retain the PPT's physical font sizes after cropping."""
    from pypdf import PdfReader

    sizes = {}
    for relative in re.findall(r"!\[[^\n]+\]\(([^)]+\.pdf)\)", SOURCE.read_text("utf-8")):
        path = (SOURCE.parent / relative).resolve()
        page = PdfReader(path).pages[0]
        sizes[path] = (float(page.mediabox.width), float(page.mediabox.height))
    return sizes


def append_methods(doc) -> None:
    """Append native editable paragraphs using the existing manuscript styles."""
    if any(p.text.strip() == "方法" for p in doc.paragraphs):
        raise ValueError("Methods already present; preserve author edits and edit in place.")
    for block in re.split(r"\n\s*\n", METHODS.read_text(encoding="utf-8").strip()):
        if block.startswith("#"):
            marks, value = block.split(" ", 1)
            p = doc.add_paragraph(value, style=f"Heading {len(marks)}")
            p.paragraph_format.keep_with_next = True
            if len(marks) == 1:
                p.paragraph_format.page_break_before = True
        else:
            p = doc.add_paragraph()
            p.paragraph_format.widow_control = True
            for token in re.split(r"(\^[^\^]+\^)", block.replace("\n", " ")):
                if token.startswith("^") and token.endswith("^"):
                    run = p.add_run(token[1:-1].replace("-", "\u2212"))
                    run.font.superscript = True
                else:
                    p.add_run(token)


def append_to_existing(path: Path) -> None:
    """Change only document.xml; retain all existing body elements and package parts."""
    from io import BytesIO

    from docx import Document
    from lxml import etree

    original = path.read_bytes()
    doc = Document(BytesIO(original))
    before = [etree.tostring(child) for child in doc.element.body]
    count = len(before) - 1  # The section properties must remain last in the body.
    assert doc.element.body[-1].tag.endswith("}sectPr")
    append_methods(doc)
    assert before[:-1] == [etree.tostring(child) for child in doc.element.body[:count]]
    assert before[-1] == etree.tostring(doc.element.body[-1])
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(original)) as source, zipfile.ZipFile(output, "w") as dest:
        dest.comment = source.comment
        for item in source.infolist():
            data = (
                doc.element.xml.encode("utf-8")
                if item.filename == "word/document.xml"
                else source.read(item.filename)
            )
            dest.writestr(item, data)
    if path.read_bytes() != original:
        raise RuntimeError("Word changed during export; preserve the newer author copy.")
    path.write_bytes(output.getvalue())
    print(f"Chinese Word: appended Methods; preserved {count} original body elements", flush=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dependency-path", type=Path)
    parser.add_argument(
        "--append-methods", action="store_true", help="Append to --output, preserving author edits"
    )
    parser.add_argument(
        "--output", type=Path, default=ROOT / "output/docx/chemworld-ncs-zh-main.docx"
    )
    parser.add_argument(
        "--build",
        type=Path,
        default=Path(tempfile.gettempdir()) / "chemworld-chinese-word",
    )
    args = parser.parse_args()
    if args.dependency_path:
        sys.path.insert(0, str(args.dependency_path))
    if args.append_methods:
        append_to_existing(args.output)
        return

    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    _, metadata, body = SOURCE.read_text(encoding="utf-8").split("---", 2)
    meta = yaml.safe_load(metadata)
    body = re.sub(r"^\\needspace\{[^\n]+\}\s*$", "", body, flags=re.M).strip()
    order = list(
        dict.fromkeys(re.findall(r"@([a-zA-Z][\w-]+)", ENGLISH.read_text(encoding="utf-8")))
    )
    numbers = {key: i + 1 for i, key in enumerate(order)}
    assert set(re.findall(r"@([a-zA-Z][\w-]+)", body)) == set(numbers)
    args.build.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.header_distance = sec.footer_distance = Cm(1.1)

    def style_font(style, size: int, east: str = "SimSun", bold: bool = False) -> None:
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key, value in (
            ("ascii", "Times New Roman"),
            ("hAnsi", "Times New Roman"),
            ("eastAsia", east),
        ):
            fonts.set(qn("w:" + key), value)
        for attr in list(fonts.attrib):
            if attr.endswith("Theme"):
                del fonts.attrib[attr]

    normal = doc.styles["Normal"]
    style_font(normal, 12)
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.widow_control = True
    for name, size, east, bold in (
        ("Title", 20, "SimHei", True),
        ("Heading 1", 16, "SimHei", True),
        ("Heading 2", 14, "SimHei", True),
        ("Caption", 11, "SimSun", False),
        ("Footer", 10, "SimSun", False),
    ):
        style_font(doc.styles[name], size, east, bold)
    for name in ("Heading 1", "Heading 2"):
        fmt = doc.styles[name].paragraph_format
        fmt.keep_with_next = True
        fmt.space_before, fmt.space_after = Pt(14), Pt(7)
    capfmt = doc.styles["Caption"].paragraph_format
    capfmt.line_spacing = 1.1
    capfmt.space_before, capfmt.space_after = Pt(5), Pt(10)
    capfmt.keep_together = True

    expected = []
    figure_sizes = figure_sizes_pt()
    # Use one physical scale for all figures, not one placement width per crop.
    figure_scale = Cm(16.6).pt / max(width for width, _ in figure_sizes.values())

    def inline(paragraph, value: str) -> str:
        """Keep citation numbers as editable superscript runs."""
        tokens = re.split(r"(\[@[^\]]+\]|\*\*.*?\*\*)", value)
        plain = []
        for token in tokens:
            if not token:
                continue
            if token.startswith("[@"):
                ids = sorted(numbers[key] for key in re.findall(r"@([a-zA-Z][\w-]+)", token))
                label = ",".join(map(str, ids))
                r = paragraph.add_run(label)
                r.font.superscript = True
                r.font.size = Pt(9)
                plain.append(label)
            elif token.startswith("**"):
                label = token[2:-2]
                paragraph.add_run(label).bold = True
                plain.append(label)
            else:
                paragraph.add_run(token)
                plain.append(token)
        return "".join(plain)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(14)
    expected.append(inline(title, meta["title"]))
    doc.add_paragraph("摘要", style="Heading 1")
    abstract = doc.add_paragraph()
    expected.append(inline(abstract, meta["abstract"].strip()))
    # Page numbers aid review; the manuscript itself remains unprotected.
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    blocks = re.split(r"\n\s*\n", body)
    # Place the tall Figure 3 just before its case paragraph, avoiding a nearly
    # empty preceding page while retaining prose order and full image width.
    figure3 = next(i for i, text in enumerate(blocks) if "/figure03.pdf)" in text)
    assert blocks[figure3 - 1].startswith("一个选定的电化学配对")
    blocks[figure3 - 1], blocks[figure3] = blocks[figure3], blocks[figure3 - 1]
    figure_count = 0
    previous_table = False
    for block in blocks:
        block = block.strip()
        if not block:
            continue
        match = re.fullmatch(r"!\[(.+)\]\(([^)]+)\)(?:\{[^}]+\})?", block, re.S)
        if match:
            figure_count += 1
            caption, rel = match.groups()
            pdf = (SOURCE.parent / rel).resolve()
            prefix = args.build / f"figure{figure_count:02d}"
            subprocess.run(
                [
                    shutil.which("pdftoppm"),
                    "-singlefile",
                    "-f",
                    "1",
                    "-l",
                    "1",
                    "-r",
                    "300",
                    "-png",
                    str(pdf),
                    str(prefix),
                ],
                check=True,
                capture_output=True,
            )
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1
            width, height = figure_sizes[pdf]
            picture = p.add_run().add_picture(
                str(prefix.with_suffix(".png")),
                width=Pt(width * figure_scale),
                height=Pt(height * figure_scale),
            )
            picture._inline.docPr.set("descr", f"图 {figure_count} {caption.split('。', 1)[0]}")
            p = doc.add_paragraph(style="Caption")
            p.add_run(f"图 {figure_count}  ").bold = True
            expected.append(f"图 {figure_count}  " + inline(p, caption))
            print(f"Chinese Word: figures {figure_count}/6", flush=True)
        elif block.startswith("#"):
            marks, text = block.split(" ", 1)
            p = doc.add_paragraph(style=f"Heading {len(marks)}")
            if text == "讨论":
                p.paragraph_format.page_break_before = True
            expected.append(inline(p, text))
        elif block.startswith("|"):
            rows = [
                [cell.strip() for cell in line.strip().strip("|").split("|")]
                for line in block.splitlines()
            ]
            rows = [row for row in rows if not all(re.fullmatch(r"[: -]+", cell) for cell in row)]
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = False
            widths = [3.2, 3.0, 3.0, 3.7, 3.7]
            for column, width in zip(table.columns, widths, strict=True):
                column.width = Cm(width)
            props = table._tbl.tblPr
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                node = OxmlElement("w:" + edge)
                for attr, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    node.set(qn("w:" + attr), value)
                borders.append(node)
            props.append(borders)
            margins = OxmlElement("w:tblCellMar")
            for edge in ("top", "left", "bottom", "right"):
                node = OxmlElement("w:" + edge)
                node.set(qn("w:w"), "100")
                node.set(qn("w:type"), "dxa")
                margins.append(node)
            props.append(margins)
            for row_index, values in enumerate(rows):
                row = table.rows[row_index]
                row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                if row_index == 0:
                    row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
                for cell, width, value in zip(row.cells, widths, values, strict=True):
                    cell.width = Cm(width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    p.alignment = (
                        WD_ALIGN_PARAGRAPH.LEFT if width == widths[0] else WD_ALIGN_PARAGRAPH.CENTER
                    )
                    p.paragraph_format.line_spacing = 1.1
                    p.paragraph_format.space_after = Pt(0)
                    expected.append(inline(p, value))
                    for run in p.runs:
                        run.font.size = Pt(11)
                        run.bold = row_index == 0
                    if row_index == 0:
                        shading = OxmlElement("w:shd")
                        shading.set(qn("w:fill"), "F2F2F2")
                        cell._tc.get_or_add_tcPr().append(shading)
        else:
            p = doc.add_paragraph()
            if previous_table:
                p.paragraph_format.space_before = Pt(7)
            if block.startswith("**表"):
                p.paragraph_format.keep_with_next = True
            expected.append(inline(p, block.replace("\n", " ")))
        previous_table = block.startswith("|")

    assert figure_count == 6 and len(doc.tables) == 1
    doc.core_properties.title = meta["title"]
    doc.core_properties.language = "zh-CN"
    doc.core_properties.author = ""
    doc.core_properties.subject = ""
    settings = doc.settings.element
    no_compress = OxmlElement("w:doNotAutoCompressPictures")
    settings.append(no_compress)
    language = normal.element.get_or_add_rPr().find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        normal.element.get_or_add_rPr().append(language)
    language.set(qn("w:eastAsia"), "zh-CN")
    language.set(qn("w:val"), "en-US")
    append_methods(doc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    check = Document(args.output)
    actual = [p.text for p in check.paragraphs]
    actual += [
        p.text
        for t in check.tables
        for row in t.rows
        for cell in row.cells
        for p in cell.paragraphs
    ]
    assert all(text in actual for text in expected), "Text coverage failure"
    assert len(check.inline_shapes) == 6 and len(check.tables) == 1
    assert not any("[@" in text or "\\needspace" in text for text in actual)
    print(
        "Chinese Word: complete; 6 figures, 1 editable table, 30 citation keys, Methods; "
        f"{args.output}",
        flush=True,
    )


if __name__ == "__main__":
    main()
